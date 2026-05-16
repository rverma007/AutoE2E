from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

# ── Sheet 1: All Test Cases ──────────────────────────────────────────────────
ws = wb.active
ws.title = "Test Cases"

headers = ["Test File", "Feature", "Test Case ID", "Test Name", "All Assertions", "Status"]
ws.append(headers)

header_fill = PatternFill("solid", start_color="1F4E79")
header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
alt_fill    = PatternFill("solid", start_color="D9E1F2")
skip_fill   = PatternFill("solid", start_color="FFE0CC")
thin        = Side(style="thin", color="AAAAAA")
border      = Border(left=thin, right=thin, top=thin, bottom=thin)

for cell in ws[1]:
    cell.font      = header_font
    cell.fill      = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border    = border
ws.row_dimensions[1].height = 22

# Each assertions string uses " | " as separator between individual asserts.
data = [
    # ── Audit Logger ────────────────────────────────────────────────────────
    (
        "test_audit_logger.py", "Audit Logger", "TC_SM_043",
        "test_audit_logger_page_loads",
        "assert loaded, 'Audit Logger page did not load. URL: {url}'",
        "Active",
    ),
    (
        "test_audit_logger.py", "Audit Logger", "TC_SM_044",
        "test_letter_actions_create_audit_entry",
        "assert audit_logger_page.is_loaded(timeout=15_000)"
        " | assert row_count > 0 or (total is not None and total > 0),"
        " 'Audit Logger shows no entries — expected at least one action to be logged.'",
        "Active",
    ),
    (
        "test_audit_logger.py", "Audit Logger", "TC_SM_045",
        "test_audit_data_correct_fields",
        "assert audit_logger_page.is_loaded(timeout=15_000)"
        " | assert len(non_empty) >= 3, 'Audit entry appears to be missing fields.'",
        "Active",
    ),
    (
        "test_audit_logger.py", "Audit Logger", "TC_SM_046",
        "test_search_and_date_range_filter",
        "assert audit_logger_page.is_loaded(timeout=15_000)"
        " | assert filtered_count >= 0, 'row_count() raised an error after filter.'"
        " | assert audit_logger_page.is_loaded(timeout=10_000),"
        " 'Audit Logger lost loaded state after applying filter.'",
        "Active",
    ),
    (
        "test_audit_logger.py", "Audit Logger", "TC_SM_047",
        "test_download_audit_report",
        "assert audit_logger_page.is_loaded(timeout=15_000)"
        " | assert filename, 'Download triggered but filename is empty.'"
        " | assert audit_logger_page.is_visible(download_button, timeout=5_000),"
        " 'Download button missing after click.'",
        "Active",
    ),
    (
        "test_audit_logger.py", "Audit Logger", "TC_SM_048",
        "test_pagination_and_rows_per_page",
        "assert audit_logger_page.is_loaded(timeout=15_000)"
        " | assert audit_logger_page.is_loaded(timeout=10_000),"
        " 'Page lost loaded state after changing rows-per-page.'"
        " | assert post_count >= 0, 'row_count() raised an error after pagination.'",
        "Active",
    ),

    # ── Component Library ───────────────────────────────────────────────────
    (
        "test_component_library.py", "Component Library", "TC_SM_028",
        "test_component_library_accessible",
        "assert loaded, 'Component Library page did not load. URL: {url}'",
        "Active",
    ),
    (
        "test_component_library.py", "Component Library", "TC_SM_029",
        "test_placeholder_list_loads",
        "assert component_library_page.is_loaded(), 'Component Library did not load.'"
        " | assert row_count >= 0, 'row_count() raised an error.'",
        "Active",
    ),
    (
        "test_component_library.py", "Component Library", "TC_SM_030",
        "test_insert_list_loads",
        "assert component_library_page.is_loaded(), 'Component Library did not load.'"
        " | assert row_count >= 0, 'row_count() raised an error.'",
        "Active",
    ),
    (
        "test_component_library.py", "Component Library", "TC_SM_031",
        "test_condition_list_loads",
        "assert component_library_page.is_loaded(), 'Component Library did not load.'"
        " | assert row_count >= 0, 'row_count() raised an error.'",
        "Active",
    ),
    (
        "test_component_library.py", "Component Library", "TC_SM_032",
        "test_block_list_loads",
        "assert component_library_page.is_loaded(), 'Component Library did not load.'"
        " | assert row_count >= 0, 'row_count() raised an error.'",
        "Active",
    ),
    (
        "test_component_library.py", "Component Library", "TC_SM_033",
        "test_user_sample_file_upload",
        "assert component_library_page.is_loaded(), 'Component Library did not load.'"
        " | assert component_library_page.is_loaded(timeout=10_000),"
        " 'Component Library page lost loaded state after upload.'",
        "Active",
    ),

    # ── Dashboard ───────────────────────────────────────────────────────────
    (
        "test_dashboard.py", "Dashboard", "N/A",
        "test_stat_cards_match_api_status_summary",
        "assert api_body, 'letter-type-version API response was not captured.'"
        " | assert summary, 'statusSummary missing from API response'"
        " | assert dashboard.is_loaded(), 'Dashboard did not reach loaded state.'"
        " | assert ui_val == api_val, 'Draft & Imported UI ≠ API'"
        " | assert ui_val == api_val, 'Pending Approvals UI ≠ API'"
        " | assert ui_val == api_val, 'Rejected Approvals UI ≠ API'"
        " | assert ui_val == api_val, 'Approved Versions UI ≠ API'",
        "Active",
    ),
    (
        "test_dashboard.py", "Dashboard", "N/A",
        "test_pending_tab_footer_matches_api",
        "assert dashboard.is_loaded()"
        " | assert card_count == api_summary, 'Pending card UI ≠ API summary.submitted'"
        " | if api_total == 0: assert footer_count == 0"
        " | else: assert footer_count == api_total, 'Pending footer ≠ API totalRecords'",
        "Active",
    ),
    (
        "test_dashboard.py", "Dashboard", "N/A",
        "test_rejected_tab_footer_matches_api",
        "assert dashboard.is_loaded()"
        " | assert card_count == expected, 'Rejected card UI ≠ statusSummary.rejected'"
        " | assert api_total == expected, 'Rejected API totalRecords ≠ statusSummary.rejected'"
        " | if api_total == 0: assert footer_count == 0"
        " | else: assert footer_count == api_total, 'Rejected footer ≠ API totalRecords'",
        "Active",
    ),
    (
        "test_dashboard.py", "Dashboard", "N/A",
        "test_approved_tab_footer_matches_api",
        "assert dashboard.is_loaded()"
        " | assert card_count == expected, 'Approved card UI ≠ statusSummary.approved'"
        " | assert api_total == expected, 'Approved API totalRecords ≠ statusSummary.approved'"
        " | if api_total == 0: assert footer_count == 0"
        " | else: assert footer_count == api_total, 'Approved footer ≠ API totalRecords'",
        "Active",
    ),

    # ── Letter Consolidation ─────────────────────────────────────────────────
    (
        "test_letter_consolidation.py", "Letter Consolidation", "TC_SM_050",
        "test_grouping_in_consolidation_page",
        "assert letter_consolidation_page.is_loaded(timeout=15_000)"
        " | assert group_count >= 0, 'group_count() raised an error.'",
        "SKIP",
    ),
    (
        "test_letter_consolidation.py", "Letter Consolidation", "TC_SM_051",
        "test_percentage_change_after_approval",
        "assert letter_consolidation_page.is_loaded(timeout=15_000)"
        " | assert letter_consolidation_page.is_loaded(timeout=5_000)  [placeholder]",
        "SKIP",
    ),
    (
        "test_letter_consolidation.py", "Letter Consolidation", "TC_SM_052",
        "test_only_one_representative_per_group",
        "assert letter_consolidation_page.is_loaded(timeout=15_000)"
        " | assert rep_visible, 'No representative badge found in Letter Consolidation.'",
        "SKIP",
    ),
    (
        "test_letter_consolidation.py", "Letter Consolidation", "TC_SM_053",
        "test_documents_can_be_moved_between_groups",
        "assert letter_consolidation_page.is_loaded(timeout=15_000)"
        " | assert letter_consolidation_page.is_loaded(timeout=5_000)  [placeholder]",
        "SKIP",
    ),

    # ── Letter Control Center ────────────────────────────────────────────────
    (
        "test_letter_control_center.py", "Letter Control Center", "TC_SM_034",
        "test_generated_letters_page_loads",
        "assert loaded, 'Letter Control Center page did not load.'"
        " | assert row_count >= 0, 'row_count() raised an error.'",
        "Active",
    ),
    (
        "test_letter_control_center.py", "Letter Control Center", "TC_SM_035",
        "test_filters_present",
        "assert letter_control_center_page.is_loaded(timeout=15_000)"
        " | assert filters_visible, 'Filter button/section not found on Letter Control Center page.'",
        "Active",
    ),
    (
        "test_letter_control_center.py", "Letter Control Center", "TC_SM_036",
        "test_xml_upload_and_generate",
        "assert letter_control_center_page.is_loaded(timeout=15_000)"
        " | assert letter_control_center_page.is_loaded(timeout=15_000),"
        " 'Letter Control Center lost loaded state after XML upload.'",
        "Active",
    ),
    (
        "test_letter_control_center.py", "Letter Control Center", "TC_SM_037",
        "test_processing_to_completed_status",
        "assert letter_control_center_page.is_loaded(timeout=15_000)"
        " | assert reached or 'completed' in final_status.lower(),"
        " 'Status did not reach Completed within 120s.'",
        "Active",
    ),
    (
        "test_letter_control_center.py", "Letter Control Center", "TC_SM_038",
        "test_pdf_and_docx_download",
        "assert letter_control_center_page.is_loaded(timeout=15_000)"
        " | assert pdf_visible, 'PDF download button disappeared during test.'",
        "Active",
    ),
    (
        "test_letter_control_center.py", "Letter Control Center", "TC_SM_039",
        "test_validation_summary",
        "assert letter_control_center_page.is_loaded(timeout=15_000)"
        " | assert letter_control_center_page.is_loaded(timeout=10_000),"
        " 'Page lost loaded state after Validation Summary click.'",
        "Active",
    ),
    (
        "test_letter_control_center.py", "Letter Control Center", "TC_SM_040",
        "test_bulk_download",
        "assert letter_control_center_page.is_loaded(timeout=15_000)"
        " | assert resp.status in (200, 201, 202),"
        " 'Bulk download API returned unexpected status'"
        " | assert letter_control_center_page.is_loaded(timeout=5_000),"
        " 'Page lost loaded state after Bulk Download click.'  [fallback if no API event]",
        "Active",
    ),

    # ── Letter Import/Export ─────────────────────────────────────────────────
    (
        "test_letter_import_export.py", "Letter Import/Export", "TC_SM_024",
        "test_import_export_page_accessible",
        "assert loaded, 'Import/Export page did not load. URL: {url}'",
        "Active",
    ),
    (
        "test_letter_import_export.py", "Letter Import/Export", "TC_SM_025",
        "test_export_generates_zip",
        "assert page.is_loaded(timeout=15_000), 'Import/Export page did not load.'"
        " | assert export_visible, 'Export button not found on Import/Export page.'"
        " | assert filename, 'Export triggered but suggested_filename is empty.'"
        " | assert filename.lower().endswith(('.zip', '.json', '.xml')),"
        " 'Unexpected export file extension'"
        " | assert page.is_visible(page.export_button, timeout=5_000),"
        " 'Export button missing after click.'  [fallback if no download event]",
        "Active",
    ),
    (
        "test_letter_import_export.py", "Letter Import/Export", "TC_SM_026",
        "test_approved_zip_import",
        "No assertions — test body is pass (placeholder, on hold pending approved ZIP artifact)",
        "SKIP",
    ),
    (
        "test_letter_import_export.py", "Letter Import/Export", "TC_SM_027",
        "test_admin_only_access_enforced",
        "assert not denied or not loaded,"
        " 'Access denied banner AND page loaded simultaneously — unexpected state.'",
        "Active",
    ),

    # ── Letter Type ──────────────────────────────────────────────────────────
    (
        "test_letter_type.py", "Letter Type", "N/A",
        "test_letter_type_list_loads",
        "assert letter_type_page.is_loaded(),"
        " 'Letter Type page did not reach loaded state — search box missing.'"
        " | assert headers, 'No column headers found in the Letter Type table.'"
        " | assert row_count >= 0, 'row_count() raised an unexpected error.'"
        " | assert total >= 0, 'Footer reported negative total'  [if total is not None]",
        "Active",
    ),
    (
        "test_letter_type.py", "Letter Type", "N/A",
        "test_configure_letter_type_upload",
        "assert letter_type_page.is_loaded(), 'Letter Type page did not load.'"
        " | assert any(s in status_after_upload.lower() for s in _INITIAL_STATES),"
        " 'Expected a recognised upload status'"
        " | assert reached_terminal,"
        " 'Status did not reach a terminal state within 60s.'",
        "Active",
    ),
    (
        "test_letter_type.py", "Letter Type", "N/A",
        "test_filter_by_status",
        "assert letter_type_page.is_loaded(), 'Letter Type page did not load.'"
        " | assert filtered_count >= 0, 'row_count() raised an error after filter.'"
        " | assert not mismatches, 'Some rows do not contain the filter value'",
        "Active",
    ),
    (
        "test_letter_type.py", "Letter Type", "N/A",
        "test_download_button_triggers_download",
        "assert letter_type_page.is_loaded(), 'Letter Type page did not load.'"
        " | assert letter_type_page.is_visible(download_button, timeout=10_000),"
        " 'Download button not found on the Letter Type listing page.'"
        " | assert filename,"
        " 'Download triggered but suggested_filename is empty.'"
        " | assert letter_type_page.is_visible(download_button, timeout=5_000),"
        " 'Download button missing after click.'  [fallback]",
        "Active",
    ),

    # ── Letter Type Details ──────────────────────────────────────────────────
    (
        "test_letter_type_details.py", "Letter Type Details", "TC_SM_013",
        "test_details_page_loads",
        "assert listing.is_loaded(), 'Letter Type listing did not load.'"
        " | assert row_count > 0, 'No rows in Letter Type listing — cannot navigate to details.'"
        " | assert loaded, 'Letter Type details page did not load.'",
        "Active",
    ),
    (
        "test_letter_type_details.py", "Letter Type Details", "TC_SM_014",
        "test_version_dropdown_visible",
        "assert listing.is_loaded()"
        " | assert listing.row_count() > 0, 'No rows to click.'"
        " | assert details.is_loaded(timeout=20_000)"
        " | assert visible, 'Version dropdown/selector not found on Letter Type details page.'",
        "Active",
    ),
    (
        "test_letter_type_details.py", "Letter Type Details", "TC_SM_015",
        "test_make_current",
        "assert listing.is_loaded()"
        " | assert listing.row_count() > 0, 'No rows to click.'"
        " | assert details.is_loaded(timeout=20_000)"
        " | assert details.is_loaded(timeout=15_000),"
        " 'Page failed to recover after Make Current action.'",
        "Active",
    ),
    (
        "test_letter_type_details.py", "Letter Type Details", "TC_SM_016",
        "test_generate_test_letter",
        "assert listing.is_loaded()"
        " | assert listing.row_count() > 0, 'No rows to click.'"
        " | assert details.is_loaded(timeout=20_000)"
        " | assert 'letter-type' in authed_page.url,"
        " 'Unexpected navigation away from letter-type context after generation.'",
        "Active",
    ),
    (
        "test_letter_type_details.py", "Letter Type Details", "TC_SM_017",
        "test_validation_summary",
        "assert listing.is_loaded()"
        " | assert listing.row_count() > 0, 'No rows to click.'"
        " | assert details.is_loaded(timeout=20_000)"
        " | assert details.is_loaded(timeout=10_000),"
        " 'Page failed to recover after Validation Summary click.'",
        "Active",
    ),

    # ── Letter Type Editor ───────────────────────────────────────────────────
    (
        "test_letter_type_editor.py", "Letter Type Editor", "TC_SM_018",
        "test_editor_loads",
        "assert listing.is_loaded()"
        " | assert details.is_loaded(timeout=20_000)"
        " | assert loaded, 'Letter editor did not render. URL: {url}'",
        "Active",
    ),
    (
        "test_letter_type_editor.py", "Letter Type Editor", "TC_SM_019",
        "test_generate_test_letter_from_editor",
        "assert listing.is_loaded() [via _open_editor]"
        " | assert listing.row_count() > 0"
        " | assert details.is_loaded(timeout=20_000)"
        " | assert 'letter-type' in authed_page.url,"
        " 'Unexpected navigation away from letter-type context.'",
        "Active",
    ),
    (
        "test_letter_type_editor.py", "Letter Type Editor", "TC_SM_020",
        "test_reset_feature",
        "assert listing.is_loaded() [via _open_editor]"
        " | assert listing.row_count() > 0"
        " | assert details.is_loaded(timeout=20_000)"
        " | assert editor.is_loaded(timeout=15_000), 'Editor did not load.'"
        " | assert editor.is_loaded(timeout=10_000),"
        " 'Editor lost loaded state after Reset click.'",
        "Active",
    ),
    (
        "test_letter_type_editor.py", "Letter Type Editor", "TC_SM_021",
        "test_diff_view",
        "assert listing.is_loaded() [via _open_editor]"
        " | assert listing.row_count() > 0"
        " | assert details.is_loaded(timeout=20_000)"
        " | assert editor.is_loaded(timeout=15_000), 'Editor did not load.'"
        " | assert diff_visible,"
        " 'Diff view container did not appear after clicking Diff View.'",
        "Active",
    ),
    (
        "test_letter_type_editor.py", "Letter Type Editor", "TC_SM_022",
        "test_generated_letter_renders",
        "assert listing.is_loaded() [via _open_editor]"
        " | assert listing.row_count() > 0"
        " | assert details.is_loaded(timeout=20_000)"
        " | assert editor.is_loaded(timeout=15_000), 'Editor did not load.'"
        " | assert preview_text,"
        " 'Letter preview pane is visible but has no text content.'",
        "Active",
    ),
    (
        "test_letter_type_editor.py", "Letter Type Editor", "TC_SM_023",
        "test_submit_for_approval",
        "assert listing.is_loaded() [via _open_editor]"
        " | assert listing.row_count() > 0"
        " | assert details.is_loaded(timeout=20_000)"
        " | assert editor.is_loaded(timeout=15_000), 'Editor did not load.'"
        " | assert 'letter-type' in authed_page.url,"
        " 'Unexpected navigation after Submit for Approval.'",
        "Active",
    ),

    # ── Navigation ───────────────────────────────────────────────────────────
    (
        "test_navigation.py", "Navigation", "N/A",
        "test_all_nav_pages_load",
        "assert loaded, 'Dashboard did not load after initial navigation.'"
        " | assert DashboardPage(authed_page).is_loaded(timeout=10_000),"
        " 'My Dashboard did not load.'"
        " | assert AskAutoPage(authed_page).is_loaded(timeout=10_000),"
        " 'Ask Auto page did not load.'"
        " | assert LetterTypePage(authed_page).is_loaded(timeout=10_000),"
        " 'Letter Type page did not load.'"
        " | assert ComponentLibraryPage(authed_page).is_loaded(timeout=10_000),"
        " 'Component Library page did not load.'"
        " | assert LetterConsolidationPage(authed_page).is_loaded(timeout=10_000),"
        " 'Letter Consolidation page did not load.'"
        " | assert LetterControlCenterPage(authed_page).is_loaded(timeout=10_000),"
        " 'Letter Control Center page did not load.'"
        " | assert SettingsPage(authed_page).is_loaded(timeout=10_000),"
        " 'Settings page did not load.'"
        " | assert AuditLoggerPage(authed_page).is_loaded(timeout=10_000),"
        " 'Audit Logger page did not load.'"
        " | assert ReconReportPage(authed_page).is_loaded(timeout=10_000),"
        " 'Recon Report page did not load.'",
        "Active",
    ),

    # ── Recon Report ─────────────────────────────────────────────────────────
    (
        "test_recon_report.py", "Recon Report", "TC_SM_049",
        "test_recon_report_page_loads",
        "assert loaded, 'Recon Report page did not load.'"
        " | assert list_visible or row_count >= 0, 'Recon Report list area is not visible.'"
        " | assert filename, 'Download triggered but filename is empty.'"
        " | assert recon_report_page.is_visible(download_button, timeout=5_000),"
        " 'Download button missing after click.'  [fallback if no download event]",
        "Active",
    ),

    # ── Sanity / Smoke ───────────────────────────────────────────────────────
    (
        "test_sanity.py", "Sanity/Smoke Tests", "N/A",
        "test_base_url_is_reachable",
        "assert login_page.is_displayed(),"
        " 'Login form not visible at BASE_URL. Is the environment up?'",
        "Active",
    ),
    (
        "test_sanity.py", "Sanity/Smoke Tests", "N/A",
        "test_login_page_has_required_fields",
        "assert login_page.is_visible(login_page.username_input), 'username field missing'"
        " | assert login_page.is_visible(login_page.password_input), 'password field missing'"
        " | assert login_page.is_visible(login_page.submit_button), 'submit button missing'",
        "Active",
    ),
    (
        "test_sanity.py", "Sanity/Smoke Tests", "N/A",
        "test_page_title_is_present",
        "assert login_page.title().strip(), 'Empty document title'",
        "Active",
    ),
    (
        "test_sanity.py", "Sanity/Smoke Tests", "N/A",
        "test_login_with_valid_credentials",
        "assert current_host == app_host,"
        " 'After login, browser is still on Keycloak host not app host.'"
        " | assert letter_type.is_loaded(timeout=60_000),"
        " 'Letter Type listing did not load after login.'",
        "Active",
    ),
    (
        "test_sanity.py", "Sanity/Smoke Tests", "N/A",
        "test_login_with_invalid_credentials_shows_error_or_blocks",
        "assert not letter_type.is_loaded(timeout=5_000),"
        " 'Landed on authenticated page with bad credentials!'",
        "Active",
    ),
    (
        "test_sanity.py", "Sanity/Smoke Tests", "N/A",
        "test_listing_loads",
        "assert letter_type_page.is_loaded(), 'Letter Type listing did not load.'",
        "Active",
    ),
    (
        "test_sanity.py", "Sanity/Smoke Tests", "N/A",
        "test_search_box_accepts_input",
        "assert letter_type_page.is_loaded()"
        " | assert letter_type_page.search_box.input_value() == 'test'",
        "Active",
    ),
    (
        "test_sanity.py", "Sanity/Smoke Tests", "N/A",
        "test_listing_has_results_or_empty_state",
        "assert letter_type_page.is_loaded()"
        " | assert rows >= 0  [trivially true; proves locator resolved without error]",
        "Active",
    ),
    (
        "test_sanity.py", "Sanity/Smoke Tests", "N/A",
        "test_authed_session_bypasses_login",
        "assert letter_type.is_loaded(),"
        " 'Stored session did not bypass login — session persistence broken.'",
        "Active",
    ),

    # ── Settings ─────────────────────────────────────────────────────────────
    (
        "test_settings.py", "Settings", "TC_SM_041",
        "test_global_configuration",
        "assert loaded, 'Settings page did not load.'"
        " | assert gc_visible,"
        " 'Global Configuration section/heading not visible on Settings page.'"
        " | assert save_visible,"
        " 'Save Settings button not found on Global Configuration page.'",
        "Active",
    ),
    (
        "test_settings.py", "Settings", "TC_SM_042",
        "test_business_unit_configuration",
        "assert settings_page.is_loaded(timeout=15_000), 'Settings page did not load.'"
        " | assert save_visible,"
        " 'Save Settings button not found on Business Unit Config page.'",
        "Active",
    ),
]

for i, row in enumerate(data, start=2):
    ws.append(list(row))
    is_skip = row[5] == "SKIP"
    fill = skip_fill if is_skip else (alt_fill if i % 2 == 0 else None)
    for cell in ws[i]:
        cell.font      = Font(name="Arial", size=10)
        cell.border    = border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if fill:
            cell.fill = fill
    ws.row_dimensions[i].height = 60

ws.freeze_panes = "A2"

col_widths = [32, 22, 14, 40, 110, 10]
for idx, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(idx)].width = w

# ── Sheet 2: Summary ─────────────────────────────────────────────────────────
ws2 = wb.create_sheet("Summary")

features = {}
for row in data:
    feat, status = row[1], row[5]
    if feat not in features:
        features[feat] = {"Active": 0, "SKIP": 0}
    features[feat][status] += 1

sum_headers = ["Feature", "Active", "SKIP", "Total"]
ws2.append(sum_headers)
for cell in ws2[1]:
    cell.font      = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    cell.fill      = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border    = border
ws2.row_dimensions[1].height = 22

for i, (feat, counts) in enumerate(sorted(features.items()), start=2):
    ws2.append([feat, counts["Active"], counts["SKIP"], f"=B{i}+C{i}"])
    fill = alt_fill if i % 2 == 0 else None
    for cell in ws2[i]:
        cell.font      = Font(name="Arial", size=10)
        cell.border    = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if fill:
            cell.fill = fill

last_data_row = 1 + len(features)
total_row     = last_data_row + 1
ws2.append([
    "TOTAL",
    f"=SUM(B2:B{last_data_row})",
    f"=SUM(C2:C{last_data_row})",
    f"=SUM(D2:D{last_data_row})",
])
for cell in ws2[total_row]:
    cell.font      = Font(name="Arial", bold=True, size=11)
    cell.fill      = PatternFill("solid", start_color="BDD7EE")
    cell.border    = border
    cell.alignment = Alignment(horizontal="center", vertical="center")

ws2.freeze_panes = "A2"
for col, w in zip(["A", "B", "C", "D"], [30, 12, 12, 12]):
    ws2.column_dimensions[col].width = w

out = r"D:\AutoPythone2e\.claude\worktrees\busy-robinson-ae5abd\TestCases_Assertions.xlsx"
wb.save(out)
print(f"Saved: {out}  ({len(data)} test cases)")
